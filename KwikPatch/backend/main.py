import os

import re

import shutil
from socket import socket

import openpyxl

import hashlib

import random

import time

import base64

import hmac

import json
import random
import smtplib
from email.mime.text import MIMEText
from dotenv import load_dotenv
_env_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), ".env")
load_dotenv(dotenv_path=_env_path, override=True)

from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials

from fastapi import Depends

import re

import shutil

import openpyxl

import openpyxl.reader.drawings

openpyxl.reader.drawings.find_images = lambda archive, path: ([], [])

from io import BytesIO

from fastapi import FastAPI, UploadFile, File, Form, Depends, HTTPException

from fastapi.middleware.cors import CORSMiddleware

from fastapi.responses import FileResponse, StreamingResponse

from fastapi.staticfiles import StaticFiles

from .database import init_db, get_db_connection, DB_PATH

from .excel_parser import parse_pi_file

from .planner import match_pi_items_to_master, process_customer_planning_sheet, safe_copy_sheet

init_db()

SECRET_KEY = os.environ.get("SECRET_KEY", "KwikPatchSuperSecretKeyChangeThisInProduction").encode()

security_bearer = HTTPBearer(auto_error=False)

def hash_password(password: str) -> str:

    salt = os.urandom(16).hex()

    hash_val = hashlib.pbkdf2_hmac('sha256', password.encode(), salt.encode(), 100000).hex()

    return f"{salt}:{hash_val}"

def verify_password(password: str, hashed: str) -> bool:

    try:

        salt, hash_val = hashed.split(":")

        check_val = hashlib.pbkdf2_hmac('sha256', password.encode(), salt.encode(), 100000).hex()

        return check_val == hash_val

    except:

        return False

def generate_token(username: str) -> str:

    payload = {

        "username": username,

        "exp": time.time() + 86400 * 7 # 7 days

    }

    payload_b64 = base64.urlsafe_b64encode(json.dumps(payload).encode()).decode()

    signature = hmac.new(SECRET_KEY, payload_b64.encode(), hashlib.sha256).hexdigest()

    return f"{payload_b64}.{signature}"

def verify_token(token: str) -> str | None:

    try:

        payload_b64, signature = token.split(".")

        expected = hmac.new(SECRET_KEY, payload_b64.encode(), hashlib.sha256).hexdigest()

        if not hmac.compare_digest(expected, signature):

            return None

        payload = json.loads(base64.urlsafe_b64decode(payload_b64.encode()).decode())

        if time.time() > payload["exp"]:

            return None

        return payload["username"]

    except:

        return None

async def get_current_user(

    credentials: HTTPAuthorizationCredentials = Depends(security_bearer),

    token: str = None

):

    tok = None

    if credentials:

        tok = credentials.credentials

    elif token:

        tok = token

    if not tok:

        raise HTTPException(status_code=401, detail="Authentication token required.")

    username = verify_token(tok)

    if not username:

        raise HTTPException(status_code=401, detail="Invalid or expired session.")

    return username

app = FastAPI(title="Kwik Patch Automated Planning System")

app.add_middleware(

    CORSMiddleware,

    allow_origins=["*"],

    allow_credentials=True,

    allow_methods=["*"],

    allow_headers=["*"],

)

from fastapi import Request

import traceback

import sys

import socket

@app.middleware("http")

async def exception_logging_middleware(request: Request, call_next):

    try:

        response = await call_next(request)

        return response

    except Exception as e:

        print("=== EXCEPTION IN API REQUEST ===", file=sys.stderr)

        traceback.print_exc(file=sys.stderr)

        print("================================", file=sys.stderr)

        raise e

PLANNING_FILE_PATH = os.environ.get("PLANNING_FILE_PATH", os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "Compound Planing file.xlsx"))

UPLOAD_DIR = os.environ.get("UPLOAD_DIR", os.path.join(os.path.dirname(os.path.abspath(__file__)), "uploads"))

os.makedirs(UPLOAD_DIR, exist_ok=True)

def send_otp_email(email: str, otp: str, purpose: str = "register"):
    smtp_host = os.environ.get("SMTP_HOST", "smtp.gmail.com")
    smtp_port = int(os.environ.get("SMTP_PORT", "587"))
    smtp_user = os.environ.get("SMTP_USER", "")
    smtp_pass = os.environ.get("SMTP_PASSWORD", "")

    subject = "KwikPatch - Email Verification Code" if purpose == "register" else "KwikPatch - Login Verification Code"
    body = (
        "Hello,\n\n"
        "Your KwikPatch verification code is:\n\n"
        f"  {otp}\n\n"
        "This code is valid for 10 minutes. Do not share it with anyone.\n\n"
        "- KwikPatch Team"
    )

    sep = "=" * 50
    print(f"\n{sep}")
    print(f"OTP EMAIL [{purpose.upper()}]")
    print(f"To     : {email}")
    print(f"Subject: {subject}")
    print(f"OTP    : {otp}")
    print(f"{sep}\n")

    if not smtp_user or not smtp_pass or smtp_pass == "your-16-char-app-password-here":
        print("SMTP not configured - OTP logged to console only.")
        return

    try:
        msg = MIMEText(body)
        msg["Subject"] = subject
        msg["From"] = smtp_user
        msg["To"] = email
        smtp_ip = socket.getaddrinfo(smtp_host, smtp_port, socket.AF_INET)[0][4][0]
        with smtplib.SMTP_SSL(smtp_ip, smtp_port, timeout=15) as server:
            server.ehlo(smtp_host)
            server.login(smtp_user, smtp_pass)
            server.send_message(msg)
        print(f"Email sent successfully to {email}")
    except Exception as ex:
        print(f"SMTP send failed: {ex}")

def log_auth_event(conn, event: str, email: str = None, mobile: str = None, status: str = "success", detail: str = None):
    try:
        conn.execute(
            "INSERT INTO auth_logs (event, email, mobile, status, detail) VALUES (?, ?, ?, ?, ?)",
            (event, email, mobile, status, detail)
        )
        conn.commit()
    except Exception as e:
        print(f"Auth log error: {e}")

def generate_and_send_otp(conn, email: str, purpose: str):
    # 1. Check rate limit (60s)
    recent_otp = conn.execute(
        "SELECT expires_at FROM otp_codes WHERE email = ? AND purpose = ? ORDER BY id DESC LIMIT 1",
        (email, purpose)
    ).fetchone()
    
    if recent_otp:
        time_since_creation = 600 - (recent_otp["expires_at"] - time.time())
        if time_since_creation < 60:
            raise HTTPException(status_code=429, detail=f"Please wait {int(60 - time_since_creation)} seconds before requesting another OTP.")
            
    # 2. Invalidate older OTPs to ensure only ONE valid OTP exists at a time
    conn.execute("DELETE FROM otp_codes WHERE email = ? AND purpose = ?", (email, purpose))
    
    # 3. Create new OTP
    otp = f"{random.randint(100000, 999999)}"
    expires_at = time.time() + 600
    conn.execute(
        "INSERT INTO otp_codes (email, otp, purpose, expires_at) VALUES (?, ?, ?, ?)",
        (email, otp, purpose, expires_at)
    )
    conn.commit()
    
    # 4. Send email
    send_otp_email(email, otp, purpose)
    return otp


@app.post("/api/auth/register")
def auth_register(payload: dict):
    username = payload.get("username", "").strip()
    email = payload.get("email", "").strip().lower()
    mobile = payload.get("mobile", "").strip()

    if not username or not email:
        raise HTTPException(status_code=400, detail="Username and email are required.")

    conn = get_db_connection()
    try:
        # Check username taken
        if conn.execute("SELECT id FROM users WHERE username = ?", (username,)).fetchone():
            raise HTTPException(status_code=400, detail="Username is already taken.")

        # Check email taken by verified user
        row = conn.execute("SELECT id, is_verified FROM users WHERE email = ?", (email,)).fetchone()
        if row:
            if row["is_verified"] == 1:
                raise HTTPException(status_code=400, detail="This email is already registered. Please sign in.")
            else:
                conn.execute("DELETE FROM users WHERE id = ?", (row["id"],))
                conn.commit()

        conn.execute(
            "INSERT INTO users (username, email, mobile, is_verified) VALUES (?, ?, ?, 0)",
            (username, email, mobile)
        )

        generate_and_send_otp(conn, email, "register")
        log_auth_event(conn, "register", email=email, mobile=mobile, status="success", detail="OTP sent")

        return {"success": True, "message": f"OTP sent to {email}. Please verify your email.", "email": email}

    except HTTPException:
        raise
    except Exception as e:
        log_auth_event(conn, "register", email=email, status="failed", detail=str(e))
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        conn.close()


@app.post("/api/auth/verify-email")
def auth_verify_email(payload: dict):
    email = payload.get("email", "").strip().lower()
    otp = payload.get("otp", "").strip()

    if not email or not otp:
        raise HTTPException(status_code=400, detail="Email and OTP are required.")

    conn = get_db_connection()
    try:
        row = conn.execute(
            "SELECT otp, expires_at FROM otp_codes WHERE email = ? AND purpose = ? ORDER BY id DESC LIMIT 1",
            (email, "register")
        ).fetchone()

        if not row or row["otp"] != otp:
            log_auth_event(conn, "verify_email", email=email, status="failed", detail="Invalid OTP")
            raise HTTPException(status_code=400, detail="Invalid verification code.")

        if time.time() > row["expires_at"]:
            log_auth_event(conn, "verify_email", email=email, status="failed", detail="OTP expired")
            raise HTTPException(status_code=400, detail="Verification code has expired. Please register again.")

        conn.execute("UPDATE users SET is_verified = 1 WHERE email = ?", (email,))
        conn.execute("DELETE FROM otp_codes WHERE email = ? AND purpose = ?", (email, "register"))
        conn.commit()

        log_auth_event(conn, "verify_email", email=email, status="success")
        return {"success": True, "message": "Email verified successfully! You can now sign in."}

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        conn.close()


@app.post("/api/auth/login")
def auth_login(payload: dict):
    email = payload.get("email", "").strip().lower()

    if not email:
        raise HTTPException(status_code=400, detail="Email is required.")

    conn = get_db_connection()
    try:
        user = conn.execute(
            "SELECT id, username, is_verified FROM users WHERE email = ?", (email,)
        ).fetchone()

        if not user:
            log_auth_event(conn, "login", email=email, status="failed", detail="Email not registered")
            raise HTTPException(status_code=401, detail="This email is not registered.")

        if user["is_verified"] == 0:
            log_auth_event(conn, "login", email=email, status="failed", detail="Email not verified")
            raise HTTPException(status_code=401, detail="Please verify your email before signing in.")

        generate_and_send_otp(conn, email, "login")
        log_auth_event(conn, "login_otp_sent", email=email, status="success")

        return {"success": True, "message": f"OTP sent to {email}. Enter it to sign in.", "email": email}

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        conn.close()


@app.post("/api/auth/verify-login")
def auth_verify_login(payload: dict):
    email = payload.get("email", "").strip().lower()
    otp = payload.get("otp", "").strip()

    if not email or not otp:
        raise HTTPException(status_code=400, detail="Email and OTP are required.")

    conn = get_db_connection()
    try:
        row = conn.execute(
            "SELECT otp, expires_at FROM otp_codes WHERE email = ? AND purpose = ? ORDER BY id DESC LIMIT 1",
            (email, "login")
        ).fetchone()

        if not row or row["otp"] != otp:
            log_auth_event(conn, "verify_login", email=email, status="failed", detail="Invalid OTP")
            raise HTTPException(status_code=400, detail="Invalid verification code.")

        if time.time() > row["expires_at"]:
            log_auth_event(conn, "verify_login", email=email, status="failed", detail="OTP expired")
            raise HTTPException(status_code=400, detail="Code has expired. Please request a new one.")

        user = conn.execute("SELECT username FROM users WHERE email = ?", (email,)).fetchone()
        conn.execute("DELETE FROM otp_codes WHERE email = ? AND purpose = ?", (email, "login"))
        conn.commit()

        token = generate_token(user["username"])
        log_auth_event(conn, "verify_login", email=email, status="success", detail=f"user: {user['username']}")

        return {"success": True, "token": token, "username": user["username"]}

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        conn.close()


@app.post("/api/auth/resend-otp")
def auth_resend_otp(payload: dict):
    email = payload.get("email", "").strip().lower()
    purpose = payload.get("purpose", "register")

    if not email:
        raise HTTPException(status_code=400, detail="Email is required.")

    conn = get_db_connection()
    try:
        generate_and_send_otp(conn, email, purpose)
        log_auth_event(conn, "resend_otp", email=email, status="success", detail=f"purpose={purpose}")

        return {"success": True, "message": f"A new OTP has been sent to {email}."}

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        conn.close()


@app.get("/api/auth/me")
def auth_me(user: str = Depends(get_current_user)):
    return {"success": True, "username": user}


@app.get("/api/customers")

def get_customers(user: str = Depends(get_current_user)):

    conn = get_db_connection()

    cursor = conn.cursor()

    cursor.execute("SELECT id, name, consignee_info FROM customers")

    rows = cursor.fetchall()

    conn.close()

    return [dict(r) for r in rows]

@app.post("/api/customers")

def create_customer(payload: dict, user: str = Depends(get_current_user)):

    name = payload.get("name")

    consignee_info = payload.get("consignee_info", "")

    if not name:

        raise HTTPException(status_code=400, detail="Customer name is required.")

    conn = get_db_connection()

    cursor = conn.cursor()

    # Check if exists

    cursor.execute("SELECT id FROM customers WHERE name = ?", (name,))

    if cursor.fetchone():

        conn.close()

        raise HTTPException(status_code=400, detail="Customer already exists.")

    try:

        # Save to DB

        cursor.execute("INSERT INTO customers (name, consignee_info) VALUES (?, ?)", (name, consignee_info))

        conn.commit()

        customer_id = cursor.lastrowid

        return {"id": customer_id, "name": name, "consignee_info": consignee_info}

    except Exception as e:

        raise HTTPException(status_code=500, detail=str(e))

    finally:

        conn.close()

@app.get("/api/customers/{customer_id}/planning")

def get_customer_planning(customer_id: int, user: str = Depends(get_current_user)):

    conn = get_db_connection()

    cursor = conn.cursor()

    # Check if customer exists

    cursor.execute("SELECT name, consignee_info FROM customers WHERE id = ?", (customer_id,))

    cust_row = cursor.fetchone()

    if not cust_row:

        conn.close()

        raise HTTPException(status_code=404, detail="Customer not found.")

    customer_name = cust_row["name"]

    consignee_info = cust_row["consignee_info"]

    # Load Master product groups

    if not os.path.exists(PLANNING_FILE_PATH):

        conn.close()

        raise HTTPException(status_code=404, detail="Planning Excel file not found.")

    try:

        wb = openpyxl.load_workbook(PLANNING_FILE_PATH, data_only=True)

        from .planner import clean_sheet_name, get_product_groups, calculate_sheet_values

        clean_tab_name = clean_sheet_name(customer_name)

        if clean_tab_name in wb.sheetnames:

            sheet = wb[clean_tab_name]

        else:

            sheet = wb["LVP + SVG"]

        groups = get_product_groups(sheet)

        # Get sum of orders for this customer from database

        cursor.execute("""

            SELECT item_name, compound, side, thou, SUM(order_qty) as total_order

            FROM planning_records

            WHERE customer_id = ?

            GROUP BY item_name, compound, side, thou

        """, (customer_id,))

        db_records = cursor.fetchall()

        # Build mapping of (item_name, compound, side, thou) -> total_order

        orders_map = {}

        for r in db_records:

            comp = str(r["compound"] or "").strip().upper()

            side = str(r["side"] or "").strip().upper()

            thou = r["thou"] or 0

            key = (str(r["item_name"]).strip().upper(), comp, side, thou)

            orders_map[key] = r["total_order"]

        # Flatten groups into a flat list of rows in order

        flat_rows = []

        for group_name, rows in groups.items():

            for row in rows:

                comp = str(row["compound"] or "").strip().upper()

                side = str(row["side"] or "").strip().upper()

                thou = row["thou"] or 0

                key = (str(row["item_name"]).strip().upper(), comp, side, thou)

                # Check if we have order quantity

                order_qty = orders_map.get(key, 0)

                # Recalculate sheets and total weight based on sum

                required_sheet = 0.0

                per_sheet_weight = 0.0

                total_kg = 0.0

                if order_qty > 0:

                    required_sheet, per_sheet_weight, total_kg = calculate_sheet_values(

                        row["comp_sheet_size"],

                        row["per_sheet_gm"],

                        row["per_sheet_item"],

                        order_qty

                    )

                # Check if this row is required (not #N/A)

                gm = row["per_sheet_gm"]

                sz = row["comp_sheet_size"]

                is_required = True

                if gm is None or sz is None:

                    is_required = False

                else:

                    gm_str = str(gm).upper().strip()

                    sz_str = str(sz).upper().strip()

                    if "N/A" in gm_str or "VALUE" in gm_str or "DIV" in gm_str or not sz_str or sz_str == "NONE":

                        is_required = False

                flat_rows.append({

                    "row_index": row["row_index"],

                    "item_name": row["item_name"],

                    "compound": row["compound"],

                    "side": row["side"],

                    "thou": row["thou"],

                    "die_type": row["die_type"],

                    "die_name": row["die_name"],

                    "die_size": row["die_size"],

                    "comp_sheet_size": row["comp_sheet_size"],

                    "per_sheet_item": row["per_sheet_item"],

                    "order_qty": order_qty if order_qty > 0 else None,

                    "required_sheet": round(required_sheet, 2) if order_qty > 0 else None,

                    "per_sheet_gm": row["per_sheet_gm"],

                    "per_sheet_weight_gm": round(per_sheet_weight, 2) if order_qty > 0 else None,

                    "total_kg": round(total_kg, 3) if order_qty > 0 else None,

                    "is_required": is_required

                })

        return {

            "customer_name": customer_name,

            "consignee_info": consignee_info,

            "rows": flat_rows

        }

    except Exception as e:

        raise HTTPException(status_code=500, detail=str(e))

    finally:

        conn.close()

@app.get("/api/reports/download-excel/{upload_id}")

def download_excel_report(upload_id: int, user: str = Depends(get_current_user)):

    conn = get_db_connection()

    cursor = conn.cursor()

    # 1. Fetch upload details

    cursor.execute("""

        SELECT uh.pi_number, uh.upload_date, c.name as customer_name, c.consignee_info

        FROM upload_history uh

        JOIN customers c ON uh.customer_id = c.id

        WHERE uh.id = ?

    """, (upload_id,))

    upload = cursor.fetchone()

    if not upload:

        conn.close()

        raise HTTPException(status_code=404, detail="Upload record not found.")

    # 2. Fetch planning records for this upload

    cursor.execute("""

        SELECT item_name, compound, side, thou, die_type, die_name, die_size, 

               comp_sheet_size, per_sheet_item, required_sheet, total_kg

        FROM planning_records

        WHERE upload_id = ?

        ORDER BY id ASC

    """, (upload_id,))

    records = cursor.fetchall()

    conn.close()

    # 3. Create Excel workbook

    import openpyxl

    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side as BorderSide

    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()

    ws = wb.active

    ws.title = "Production Compound Requirement"

    # Show gridlines explicitly

    ws.views.sheetView[0].showGridLines = True

    # Font styles

    title_font = Font(name="Calibri", size=16, bold=True, color="0F172A")

    meta_font = Font(name="Calibri", size=11, bold=True, color="334155")

    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")

    data_font = Font(name="Calibri", size=10, color="000000")

    bold_data_font = Font(name="Calibri", size=10, bold=True, color="000000")

    total_font = Font(name="Calibri", size=12, bold=True, color="16A34A")

    # Alignments

    center_align = Alignment(horizontal="center", vertical="center")

    left_align = Alignment(horizontal="left", vertical="center")

    right_align = Alignment(horizontal="right", vertical="center")

    # Fills

    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")

    group_header_fill = PatternFill(start_color="CBD5E1", end_color="CBD5E1", fill_type="solid")

    subtotal_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

    alt_row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    # Borders

    thin_side = BorderSide(border_style="thin", color="E2E8F0")

    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    # Write Title

    ws.merge_cells("A2:K2")

    ws["A2"] = "KWIKPATCH COMPOUND REQUIREMENT REPORT"

    ws["A2"].font = title_font

    ws["A2"].alignment = center_align

    # Write Meta Details

    ws.merge_cells("A3:K3")

    ws["A3"] = f"Customer Name: {upload['customer_name']}"

    ws["A3"].font = meta_font

    ws["A3"].alignment = center_align

    ws.merge_cells("A4:K4")

    ws["A4"] = f"PI Number: {upload['pi_number']}"

    ws["A4"].font = meta_font

    ws["A4"].alignment = center_align

    ws.merge_cells("A5:K5")

    ws["A5"] = f"Date: {upload['upload_date']}"

    ws["A5"].font = meta_font

    ws["A5"].alignment = center_align

    # Header Row

    headers = [

        "Item Name", "Compound", "Side", "Thickness (Thou)", "Die Type", 

        "Die Name", "Die Size (MM)", "Comp Sheet Size", "Per Sheet Item", 

        "Req. Sheets", "Total kg"

    ]

    header_row_idx = 7

    ws.row_dimensions[header_row_idx].height = 28

    for col_idx, h_text in enumerate(headers, 1):

        cell = ws.cell(row=header_row_idx, column=col_idx, value=h_text)

        cell.font = header_font

        cell.fill = header_fill

        cell.alignment = center_align

        cell.border = thin_border

    # Group by compound

    from collections import defaultdict

    groups = defaultdict(list)

    for r in records:

        comp = str(r["compound"] or "-").strip().upper()

        groups[comp].append(r)

    current_row_idx = header_row_idx + 1

    compound_totals = {}

    grand_total_kg = 0.0

    for compound, group_records in groups.items():

        comp_total_kg = 0.0

        # 1. Compound Group Section Header

        ws.merge_cells(start_row=current_row_idx, start_column=1, end_row=current_row_idx, end_column=11)

        cell = ws.cell(row=current_row_idx, column=1, value=f"Compound Group: {compound}")

        cell.font = bold_data_font

        cell.alignment = left_align

        ws.row_dimensions[current_row_idx].height = 22

        # Apply border & fill to all merged cells in this row

        for c in range(1, 12):

            cell_m = ws.cell(row=current_row_idx, column=c)

            cell_m.fill = group_header_fill

            cell_m.border = thin_border

        current_row_idx += 1

        # 2. Group Rows

        for r_idx, r in enumerate(group_records):

            kg = r["total_kg"] or 0.0

            comp_total_kg += kg

            grand_total_kg += kg

            row_vals = [

                r["item_name"],

                r["compound"] or "-",

                r["side"] or "-",

                r["thou"] or "-",

                r["die_type"] or "-",

                r["die_name"] or "-",

                r["die_size"] or "-",

                r["comp_sheet_size"] or "-",

                r["per_sheet_item"] or "-",

                round(r["required_sheet"], 2) if r["required_sheet"] is not None else "-",

                round(kg, 3) if kg > 0 else "-"

            ]

            ws.row_dimensions[current_row_idx].height = 20

            bg_fill = alt_row_fill if r_idx % 2 == 0 else white_fill

            for col_idx, val in enumerate(row_vals, 1):

                cell = ws.cell(row=current_row_idx, column=col_idx, value=val)

                cell.font = bold_data_font if col_idx in [1, 11] else data_font

                cell.fill = bg_fill

                cell.alignment = center_align

                cell.border = thin_border

            current_row_idx += 1

        compound_totals[compound] = comp_total_kg

        # 3. Group Subtotal Row

        ws.merge_cells(start_row=current_row_idx, start_column=1, end_row=current_row_idx, end_column=10)

        sub_label_cell = ws.cell(row=current_row_idx, column=1, value=f"Total for {compound}:")

        sub_label_cell.font = bold_data_font

        sub_label_cell.alignment = left_align

        val_cell = ws.cell(row=current_row_idx, column=11, value=round(comp_total_kg, 3))

        val_cell.font = bold_data_font

        val_cell.alignment = center_align

        ws.row_dimensions[current_row_idx].height = 22

        for c in range(1, 12):

            cell_m = ws.cell(row=current_row_idx, column=c)

            cell_m.fill = subtotal_fill

            cell_m.border = thin_border

        current_row_idx += 1

    # Spacer row

    current_row_idx += 1

    # 4. Summary Totals at the bottom

    for comp, weight in compound_totals.items():

        ws.merge_cells(start_row=current_row_idx, start_column=1, end_row=current_row_idx, end_column=11)

        cell = ws.cell(row=current_row_idx, column=1, value=f"{comp} TOTAL REQUIRED WEIGHT: {weight:.3f} kg")

        cell.font = total_font

        cell.alignment = center_align

        ws.row_dimensions[current_row_idx].height = 22

        current_row_idx += 1

    current_row_idx += 1

    ws.merge_cells(start_row=current_row_idx, start_column=1, end_row=current_row_idx, end_column=11)

    cell = ws.cell(row=current_row_idx, column=1, value=f"GRAND TOTAL REQUIRED WEIGHT: {grand_total_kg:.3f} kg")

    cell.font = total_font

    cell.alignment = center_align

    ws.row_dimensions[current_row_idx].height = 24

    # Auto-adjust column widths

    for col in ws.columns:

        max_len = 0

        col_letter = get_column_letter(col[0].column)

        for cell in col:

            if cell.row < header_row_idx:

                continue

            if cell.value is not None:

                max_len = max(max_len, len(str(cell.value)))

        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    import io

    excel_buffer = io.BytesIO()

    wb.save(excel_buffer)

    excel_buffer.seek(0)

    filename = f"Production_Summary_{upload['customer_name']}_{upload['pi_number']}.xlsx".replace(" ", "_")

    from fastapi.responses import StreamingResponse

    return StreamingResponse(

        excel_buffer,

        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",

        headers={"Content-Disposition": f"attachment; filename={filename}"}

    )

@app.get("/api/reports/download-pdf/{upload_id}")

def download_pdf_receipt(upload_id: int, user: str = Depends(get_current_user)):

    conn = get_db_connection()

    cursor = conn.cursor()

    # 1. Fetch upload history details

    cursor.execute("""

        SELECT uh.pi_number, uh.upload_date, c.name as customer_name, c.consignee_info

        FROM upload_history uh

        JOIN customers c ON uh.customer_id = c.id

        WHERE uh.id = ?

    """, (upload_id,))

    upload = cursor.fetchone()

    if not upload:

        conn.close()

        raise HTTPException(status_code=404, detail="Upload record not found.")

    # 2. Fetch planning records for this upload

    cursor.execute("""

        SELECT item_name, compound, side, thou, die_type, die_name, die_size, 

               comp_sheet_size, per_sheet_item, required_sheet, total_kg

        FROM planning_records

        WHERE upload_id = ?

        ORDER BY id ASC

    """, (upload_id,))

    records = cursor.fetchall()

    conn.close()

    # 3. Generate ReportLab PDF

    from reportlab.lib.pagesizes import A4, landscape

    from reportlab.lib import colors

    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether

    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    import io

    pdf_buffer = io.BytesIO()

    # Create document in landscape A4 to fit 11 columns beautifully

    doc = SimpleDocTemplate(

        pdf_buffer,

        pagesize=landscape(A4),

        leftMargin=20,

        rightMargin=20,

        topMargin=20,

        bottomMargin=20

    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(

        'ReceiptTitle',

        parent=styles['Heading1'],

        fontName='Helvetica-Bold',

        fontSize=20,

        leading=24,

        textColor=colors.HexColor('#0f172a'),

        alignment=1  # Center

    )

    meta_style = ParagraphStyle(

        'ReceiptMeta',

        fontName='Helvetica-Bold',

        fontSize=12,

        leading=16,

        textColor=colors.HexColor('#334155'),

        alignment=1  # Center

    )

    table_header_style = ParagraphStyle(

        'TableHeader',

        fontName='Helvetica-Bold',

        fontSize=9,

        leading=11,

        textColor=colors.white,

        alignment=1 # Center

    )

    table_cell_style = ParagraphStyle(

        'TableCell',

        fontName='Helvetica',

        fontSize=8,

        leading=10,

        textColor=colors.HexColor('#1e293b'),

        alignment=1 # Center

    )

    table_cell_bold_style = ParagraphStyle(

        'TableCellBold',

        fontName='Helvetica-Bold',

        fontSize=8,

        leading=10,

        textColor=colors.HexColor('#0f172a'),

        alignment=1 # Center

    )

    story = []

    # Header block - Customer Name, PI No, Date (Center aligned at the top of the page)

    story.append(Paragraph("KWIKPATCH COMPOUND REQUIREMENT RECEIPT", title_style))

    story.append(Spacer(1, 10))

    story.append(Paragraph(f"Customer Name: {upload['customer_name']}", meta_style))

    story.append(Paragraph(f"PI Number: {upload['pi_number']}", meta_style))

    story.append(Paragraph(f"Date: {upload['upload_date']}", meta_style))

    story.append(Spacer(1, 15))

    # Table headers

    table_data = [[

        Paragraph("Item Name", table_header_style),

        Paragraph("Compound", table_header_style),

        Paragraph("Side", table_header_style),

        Paragraph("Thickness (Thou)", table_header_style),

        Paragraph("Die Type", table_header_style),

        Paragraph("Die Name", table_header_style),

        Paragraph("Die Size (MM)", table_header_style),

        Paragraph("Comp Sheet Size", table_header_style),

        Paragraph("Per Sheet Item", table_header_style),

        Paragraph("Req. Sheets", table_header_style),

        Paragraph("Total kg", table_header_style)

    ]]

    # Group by compound to display same compound together

    from collections import defaultdict

    groups = defaultdict(list)

    for r in records:

        comp = str(r["compound"] or "-").strip().upper()

        groups[comp].append(r)

    total_kg_sum = 0.0

    t_style_cmds = [

        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0f172a')),

        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),

        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),

        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),

        ('TOPPADDING', (0, 0), (-1, -1), 5),

        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),

    ]

    current_row_idx = 1

    compound_totals = {}

    # Iterate through compound groups

    for compound, group_records in groups.items():

        comp_total_kg = 0.0

        # Add a section header row for this compound group

        table_data.append([

            Paragraph(f"<b>Compound Group: {compound}</b>", ParagraphStyle('CompGrp', fontName='Helvetica-Bold', fontSize=9, leading=11, textColor=colors.HexColor('#0f172a'))),

            None, None, None, None, None, None, None, None, None, None

        ])

        t_style_cmds.append(('SPAN', (0, current_row_idx), (-1, current_row_idx)))

        t_style_cmds.append(('BACKGROUND', (0, current_row_idx), (-1, current_row_idx), colors.HexColor('#cbd5e1')))

        current_row_idx += 1

        # Add rows for this compound

        for r_idx, r in enumerate(group_records):

            kg = r["total_kg"] or 0.0

            comp_total_kg += kg

            total_kg_sum += kg

            table_data.append([

                Paragraph(str(r["item_name"]), table_cell_bold_style),

                Paragraph(str(r["compound"] or "-"), table_cell_style),

                Paragraph(str(r["side"] or "-"), table_cell_style),

                Paragraph(str(r["thou"] or "-"), table_cell_style),

                Paragraph(str(r["die_type"] or "-"), table_cell_style),

                Paragraph(str(r["die_name"] or "-"), table_cell_style),

                Paragraph(str(r["die_size"] or "-"), table_cell_style),

                Paragraph(str(r["comp_sheet_size"] or "-"), table_cell_style),

                Paragraph(str(r["per_sheet_item"] or "-"), table_cell_style),

                Paragraph(f"{r['required_sheet']:.2f}" if r["required_sheet"] is not None else "-", table_cell_style),

                Paragraph(f"{kg:.3f}" if kg > 0 else "-", table_cell_bold_style)

            ])

            bg_color = colors.HexColor('#f8fafc') if r_idx % 2 == 0 else colors.white

            t_style_cmds.append(('BACKGROUND', (0, current_row_idx), (-1, current_row_idx), bg_color))

            current_row_idx += 1

        compound_totals[compound] = comp_total_kg

        # Add a summary row for this compound group

        table_data.append([

            Paragraph(f"<b>Total for {compound}:</b>", table_cell_bold_style),

            None, None, None, None, None, None, None, None, None,

            Paragraph(f"<b>{comp_total_kg:.3f} kg</b>", table_cell_bold_style)

        ])

        t_style_cmds.append(('SPAN', (0, current_row_idx), (9, current_row_idx)))

        t_style_cmds.append(('BACKGROUND', (0, current_row_idx), (-1, current_row_idx), colors.HexColor('#e2e8f0')))

        t_style_cmds.append(('ALIGN', (0, current_row_idx), (0, current_row_idx), 'LEFT'))

        current_row_idx += 1

    col_widths = [80, 70, 60, 70, 70, 90, 75, 75, 65, 75, 70]

    t = Table(table_data, colWidths=col_widths, repeatRows=1)

    t_style = TableStyle(t_style_cmds)

    t.setStyle(t_style)

    story.append(t)

    story.append(Spacer(1, 15))

    # Centered total summary at the bottom showing separate weights for each compound

    total_style = ParagraphStyle(

        'TotalSummary',

        fontName='Helvetica-Bold',

        fontSize=12,

        leading=16,

        textColor=colors.HexColor('#16a34a'),

        alignment=1  # Center

    )

    for comp, weight in compound_totals.items():

        story.append(Paragraph(f"{comp} TOTAL REQUIRED WEIGHT: {weight:.3f} kg", total_style))

        story.append(Spacer(1, 4))

    story.append(Spacer(1, 10))

    grand_total_style = ParagraphStyle(

        'GrandTotalSummary',

        fontName='Helvetica-Bold',

        fontSize=14,

        leading=18,

        textColor=colors.HexColor('#16a34a'),

        alignment=1  # Center

    )

    story.append(Paragraph(f"GRAND TOTAL REQUIRED WEIGHT: {total_kg_sum:.3f} kg", grand_total_style))

    doc.build(story)

    pdf_buffer.seek(0)

    filename = f"Production_Summary_{upload['customer_name']}_{upload['pi_number']}.pdf".replace(" ", "_")

    from fastapi.responses import StreamingResponse

    return StreamingResponse(

        pdf_buffer,

        media_type="application/pdf",

        headers={"Content-Disposition": f"attachment; filename={filename}"}

    )
@app.post("/api/customers/{customer_id}/planning/save")
def save_customer_planning(customer_id: int, payload: dict, user: str = Depends(get_current_user)):

    rows = payload.get("rows", [])

    if not rows:

        raise HTTPException(status_code=400, detail="No row data provided.")

    conn = get_db_connection()

    cursor = conn.cursor()

    # Get customer details

    cursor.execute("SELECT name FROM customers WHERE id = ?", (customer_id,))

    cust_row = cursor.fetchone()

    if not cust_row:

        conn.close()

        raise HTTPException(status_code=404, detail="Customer not found.")

    customer_name = cust_row["name"]

    conn.close()

    try:

        # Open Excel for editing

        wb = openpyxl.load_workbook(PLANNING_FILE_PATH, data_only=False)

        from .planner import clean_sheet_name, set_cell_value, calculate_sheet_values

        clean_tab_name = clean_sheet_name(customer_name)

        if clean_tab_name not in wb.sheetnames:

            raise HTTPException(status_code=404, detail=f"Worksheet tab '{clean_tab_name}' not found in workbook.")

        sheet = wb[clean_tab_name]

        for r_data in rows:

            r = r_data.get("row_index")

            if not r:

                continue

            # Write back simple columns

            item_name = r_data.get("item_name")

            if item_name:

                set_cell_value(sheet, r, 2, item_name)

            set_cell_value(sheet, r, 3, r_data.get("compound"))

            set_cell_value(sheet, r, 4, r_data.get("side"))

            try:

                thou = int(r_data.get("thou")) if r_data.get("thou") is not None else None

            except:

                thou = r_data.get("thou")

            set_cell_value(sheet, r, 5, thou)

            set_cell_value(sheet, r, 6, r_data.get("die_type"))

            set_cell_value(sheet, r, 7, r_data.get("die_name"))

            set_cell_value(sheet, r, 8, r_data.get("die_size"))

            set_cell_value(sheet, r, 9, r_data.get("comp_sheet_size"))

            try:

                per_sheet_item = int(r_data.get("per_sheet_item")) if r_data.get("per_sheet_item") is not None else None

            except:

                per_sheet_item = r_data.get("per_sheet_item")

            set_cell_value(sheet, r, 10, per_sheet_item)

            per_sheet_gm = r_data.get("per_sheet_gm")

            set_cell_value(sheet, r, 13, per_sheet_gm)

            # Check if this row is required

            is_required = True

            if per_sheet_gm is None or r_data.get("comp_sheet_size") is None:

                is_required = False

            else:

                gm_str = str(per_sheet_gm).upper().strip()

                sz_str = str(r_data.get("comp_sheet_size")).upper().strip()

                if "N/A" in gm_str or "VALUE" in gm_str or "DIV" in gm_str or not sz_str or sz_str == "NONE":

                    is_required = False

            # Recalculate if it has order quantity

            order_qty = r_data.get("order_qty")

            if order_qty and is_required:

                set_cell_value(sheet, r, 11, order_qty)

                set_cell_value(sheet, r, 12, f"=K{r}/J{r}")

                set_cell_value(sheet, r, 14, f'=LEFT(I{r},FIND("X",I{r})-1)*RIGHT(I{r},LEN(I{r})-FIND("X",I{r})-1)*M{r}')

                set_cell_value(sheet, r, 15, f"=(N{r}*L{r})/1000")

            else:

                set_cell_value(sheet, r, 11, None)

                set_cell_value(sheet, r, 12, None)

                set_cell_value(sheet, r, 14, None)

                set_cell_value(sheet, r, 15, None)

        wb.save(PLANNING_FILE_PATH)

        return {"success": True, "message": "Planning sheet updated successfully."}

    except Exception as e:

        raise HTTPException(status_code=500, detail=str(e))

@app.get("/api/customers/{customer_id}/history")

def get_customer_history(customer_id: int, user: str = Depends(get_current_user)):

    conn = get_db_connection()

    cursor = conn.cursor()

    # Check if customer exists

    cursor.execute("SELECT name FROM customers WHERE id = ?", (customer_id,))

    cust = cursor.fetchone()

    if not cust:

        conn.close()

        raise HTTPException(status_code=404, detail="Customer not found.")

    customer_name = cust["name"]

    cursor.execute("""

        SELECT uh.upload_date, uh.pi_number, pr.item_name, pr.compound, pr.side, pr.order_qty, pr.total_kg

        FROM planning_records pr

        JOIN upload_history uh ON pr.upload_id = uh.id

        WHERE pr.customer_id = ?

        ORDER BY uh.upload_date DESC, pr.id ASC

    """, (customer_id,))

    rows = cursor.fetchall()

    conn.close()

    history_rows = []

    for r in rows:

        history_rows.append({

            "upload_date": r["upload_date"],

            "pi_number": r["pi_number"],

            "item_name": r["item_name"],

            "compound": r["compound"] or "",

            "side": r["side"] or "-",

            "order_qty": r["order_qty"],

            "total_kg": round(r["total_kg"] or 0.0, 3)

        })

    return {

        "customer_name": customer_name,

        "rows": history_rows

    }

@app.get("/api/master-items")

def get_master_items(customer_id: int = None, user: str = Depends(get_current_user)):

    if not os.path.exists(PLANNING_FILE_PATH):

        raise HTTPException(status_code=404, detail="Compound Planing file.xlsx not found.")

    try:

        wb = openpyxl.load_workbook(PLANNING_FILE_PATH, data_only=True)

        sheet = wb["LVP + SVG"]

        if customer_id:

            conn = get_db_connection()

            cursor = conn.cursor()

            cursor.execute("SELECT name FROM customers WHERE id = ?", (customer_id,))

            row = cursor.fetchone()

            conn.close()

            if row:

                from .planner import clean_sheet_name

                clean_tab_name = clean_sheet_name(row["name"])

                if clean_tab_name in wb.sheetnames:

                    sheet = wb[clean_tab_name]

        from .planner import get_product_groups

        groups = get_product_groups(sheet)

        return sorted(list(groups.keys()))

    except Exception as e:

        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/upload-pi")

async def upload_pi(file: UploadFile = File(...), user: str = Depends(get_current_user)):

    await file.seek(0)

    file_path = os.path.join(UPLOAD_DIR, file.filename)

    with open(file_path, "wb") as buffer:

        shutil.copyfileobj(file.file, buffer)

    conn = get_db_connection()

    cursor = conn.cursor()

    try:

        parsed_results = parse_pi_file(file_path)

        if not parsed_results:

            raise HTTPException(status_code=400, detail="No yellow highlighted items found.")

        response_data = []

        for result in parsed_results:

            cust_name = result["customer_name"]

            # Fetch all customers in the database for mapping

            cursor.execute("SELECT id, name FROM customers")

            db_customers = cursor.fetchall()

            # Find a match using prefix and substring comparison

            matched_cust = None

            parsed_clean = cust_name.strip().upper()

            for c in db_customers:

                db_clean = c["name"].strip().upper()

                db_prefix = re.split(r'\s+PI\s*|\s*PI\s+|\s+INVOICE', db_clean, flags=re.IGNORECASE)[0].strip()

                if db_prefix == parsed_clean or db_clean.startswith(parsed_clean) or parsed_clean.startswith(db_clean):

                    matched_cust = c

                    break

            if matched_cust:

                customer_id = matched_cust["id"]

                cust_name = matched_cust["name"]

            else:

                # Dynamically create customer and duplicate "LVP + SVG" template in planning sheet

                consignee = result.get("consignee_info") or f"Customer {cust_name} automatically created"

                cursor.execute(

                    "INSERT INTO customers (name, consignee_info) VALUES (?, ?)",

                    (cust_name, consignee)

                )

                conn.commit()

                customer_id = cursor.lastrowid

                # Load planning excel and duplicate template tab

                wb_plan = openpyxl.load_workbook(PLANNING_FILE_PATH, data_only=False)

                from .planner import clean_sheet_name

                clean_name = clean_sheet_name(cust_name)

                if clean_name not in wb_plan.sheetnames:

                    safe_copy_sheet(wb_plan, "LVP + SVG", clean_name)

                    wb_plan.save(PLANNING_FILE_PATH)

                    print(f"Created new tab '{clean_name}' for customer '{cust_name}' in master file.")

            pi_num = ""

            pi_match = re.search(r'PI\s*(\d+)', result["sheet_name"], re.IGNORECASE)

            if pi_match:

                pi_num = f"PI {pi_match.group(1)}"

            else:

                pi_num = result["sheet_name"]

            cursor.execute(

                "INSERT INTO upload_history (filename, customer_id, pi_number, status) VALUES (?, ?, ?, ?)",

                (file.filename, customer_id, pi_num, "pending_mapping")

            )

            conn.commit()

            upload_id = cursor.lastrowid

            matched_items, master_keys = match_pi_items_to_master(customer_id, result["items"], conn, PLANNING_FILE_PATH)

            response_data.append({

                "upload_id": upload_id,

                "customer_id": customer_id,

                "customer_name": cust_name,

                "consignee_info": result["consignee_info"],

                "pi_number": pi_num,

                "sheet_name": result["sheet_name"],

                "items": matched_items,

                "master_keys": master_keys

            })

        return {"success": True, "uploads": response_data}

    except HTTPException:

        raise

    except Exception as e:

        sz_disk = os.path.getsize(file_path) if os.path.exists(file_path) else -1

        raise HTTPException(status_code=500, detail=f"File: {file_path}, size: {sz_disk} bytes, error: {str(e)}")

    finally:

        conn.close()

@app.post("/api/upload-master-workbook")

async def upload_master_workbook(file: UploadFile = File(...), user: str = Depends(get_current_user)):

    if not file.filename.endswith(".xlsx"):

        raise HTTPException(status_code=400, detail="Only .xlsx Excel files are supported.")

    try:

        # Save file to planning path

        with open(PLANNING_FILE_PATH, "wb") as buffer:

            shutil.copyfileobj(file.file, buffer)

        # Discover customers from worksheets

        wb = openpyxl.load_workbook(PLANNING_FILE_PATH, data_only=True)

        conn = get_db_connection()

        cursor = conn.cursor()

        # Clear old customers to keep database in sync with uploaded sheet

        cursor.execute("DELETE FROM customers")

        imported_customers = []

        for name in wb.sheetnames:

            if name in ["Final output like this", "LVP + SVG", "Production Compound Requirement", "Customer Breakdown"] or name.startswith("Sheet"):

                continue

            cursor.execute("INSERT INTO customers (name, consignee_info) VALUES (?, ?)", 

                           (name, f"Customer {name} imported from master planning sheet"))

            imported_customers.append(name)

        conn.commit()

        conn.close()

        return {

            "success": True,

            "message": "Master Planning File uploaded successfully!",

            "imported_customers": imported_customers

        }

    except Exception as e:

        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/process-planning")

async def process_planning(payload: dict, user: str = Depends(get_current_user)):

    customer_name = payload.get("customer_name")

    upload_id = payload.get("upload_id")

    confirmed_mappings = payload.get("confirmed_mappings")

    if not customer_name or not upload_id or confirmed_mappings is None:

        raise HTTPException(status_code=400, detail="Missing required parameters.")

    conn = get_db_connection()

    try:

        records_count = process_customer_planning_sheet(

            customer_name, confirmed_mappings, PLANNING_FILE_PATH, conn, upload_id

        )

        cursor = conn.cursor()

        cursor.execute("UPDATE upload_history SET status = 'success' WHERE id = ?", (upload_id,))

        # Fetch the newly created planning records for display

        cursor.execute("""

            SELECT item_name, compound, side, order_qty, total_kg

            FROM planning_records

            WHERE upload_id = ?

        """, (upload_id,))

        rows = cursor.fetchall()

        conn.commit()

        processed_rows = []

        for r in rows:

            processed_rows.append({

                "item_name": r["item_name"],

                "compound": r["compound"] or "",

                "side": r["side"] or "-",

                "order_qty": r["order_qty"],

                "total_kg": round(r["total_kg"] or 0.0, 3)

            })

        return {

            "success": True,

            "upload_id": upload_id,

            "customer_name": customer_name,

            "records_processed": records_count,

            "rows": processed_rows

        }

    except Exception as e:

        raise HTTPException(status_code=500, detail=str(e))

    finally:

        conn.close()

@app.get("/api/reports")

def get_reports(user: str = Depends(get_current_user)):

    conn = get_db_connection()

    cursor = conn.cursor()

    cursor.execute("""

        SELECT uh.id, uh.filename, c.name as customer_name, uh.pi_number, uh.upload_date,

               (SELECT SUM(total_kg) FROM planning_records WHERE upload_id = uh.id) as total_kg

        FROM upload_history uh

        JOIN customers c ON uh.customer_id = c.id

        WHERE uh.status = 'success'

    """)

    rows = cursor.fetchall()

    conn.close()

    results = []

    for r in rows:

        results.append({

            "id": r["id"],

            "filename": r["filename"],

            "customer_name": r["customer_name"],

            "pi_number": r["pi_number"],

            "upload_date": r["upload_date"],

            "total_kg": round(r["total_kg"] or 0.0, 2)

        })

    return results

@app.get("/api/reports/download-workbook")

def download_workbook(user: str = Depends(get_current_user)):

    if not os.path.exists(PLANNING_FILE_PATH):

        raise HTTPException(status_code=404, detail="Compound Planing file.xlsx not found.")

    return FileResponse(

        PLANNING_FILE_PATH,

        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",

        filename="Compound Planing file.xlsx"

    )

@app.get("/api/reports/download-consolidated")

def download_consolidated(user: str = Depends(get_current_user)):

    conn = get_db_connection()

    cursor = conn.cursor()

    cursor.execute("""

        SELECT c.name as customer_name, c.consignee_info, uh.pi_number, uh.upload_date,

               pr.item_name, pr.compound, pr.side, pr.thou, pr.die_type, pr.die_name,

               pr.die_size, pr.comp_sheet_size, pr.per_sheet_item, pr.required_sheet, pr.total_kg

        FROM planning_records pr

        JOIN customers c ON pr.customer_id = c.id

        JOIN upload_history uh ON pr.upload_id = uh.id

        ORDER BY c.name ASC, uh.upload_date DESC, pr.id ASC

    """)

    records = cursor.fetchall()

    conn.close()

    if not records:

        raise HTTPException(status_code=400, detail="No planning records available.")

    from collections import defaultdict

    customer_groups = defaultdict(list)

    for r in records:

        key = (r["customer_name"], r["consignee_info"] or "", r["pi_number"] or "", r["upload_date"] or "")

        customer_groups[key].append(r)

    wb = openpyxl.Workbook()

    ws = wb.active

    ws.title = "Production Compound Requirement"

    from openpyxl.styles import Font, PatternFill, Alignment

    title_font = Font(name="Calibri", size=11, bold=True)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")

    current_row = 1

    for (cust_name, address, pi_num, date_str), items in customer_groups.items():

        # Write Customer Info block

        ws.cell(row=current_row, column=1, value="Customer Name:").font = title_font

        ws.cell(row=current_row, column=2, value=cust_name).font = Font(name="Calibri", size=11)

        current_row += 1

        ws.cell(row=current_row, column=1, value="Consignee Address:").font = title_font

        ws.cell(row=current_row, column=2, value=address).font = Font(name="Calibri", size=11)

        ws.cell(row=current_row, column=2).alignment = Alignment(wrap_text=True)

        current_row += 1

        ws.cell(row=current_row, column=1, value="PI Number:").font = title_font

        ws.cell(row=current_row, column=2, value=pi_num).font = Font(name="Calibri", size=11)

        current_row += 1

        ws.cell(row=current_row, column=1, value="Date:").font = title_font

        ws.cell(row=current_row, column=2, value=date_str).font = Font(name="Calibri", size=11)

        current_row += 2 # Leave a row blank

        # Write Table Headers

        headers = [

            "Item Name", "Compound", "Side", "Thickness (Thou)", 

            "Die Type", "Die Name", "Die Size", "Comp Sheet Size in Inch", 

            "Per Sheet Item", "Required Sheet", "Total kg"

        ]

        for col_idx, h_text in enumerate(headers, 1):

            cell = ws.cell(row=current_row, column=col_idx, value=h_text)

            cell.font = header_font

            cell.fill = header_fill

            cell.alignment = Alignment(horizontal="center")

        current_row += 1

        table_start_row = current_row

        # Write Items

        for item in items:

            row_vals = [

                item["item_name"],

                item["compound"] or "",

                item["side"] or "-",

                item["thou"] or "",

                item["die_type"] or "",

                item["die_name"] or "",

                item["die_size"] or "",

                item["comp_sheet_size"] or "",

                item["per_sheet_item"] or "",

                round(item["required_sheet"] or 0.0, 2) if item["required_sheet"] else "",

                round(item["total_kg"] or 0.0, 3) if item["total_kg"] else 0.0

            ]

            for col_idx, val in enumerate(row_vals, 1):

                cell = ws.cell(row=current_row, column=col_idx, value=val)

                if col_idx in [4, 9, 10, 11]: # Numeric alignments

                    cell.alignment = Alignment(horizontal="right")

            current_row += 1

        # Write Customer Total row

        ws.cell(row=current_row, column=1, value=f"Total for {cust_name}:").font = title_font

        ws.cell(row=current_row, column=11, value=f"=SUM(K{table_start_row}:K{current_row-1})").font = title_font

        ws.cell(row=current_row, column=11).alignment = Alignment(horizontal="right")

        current_row += 3 # Leave empty rows before next customer

    # Auto-adjust column widths

    for col in ws.columns:

        max_len = 0

        for cell in col:

            val_str = str(cell.value or '')

            if "\n" in val_str:

                val_str = max(val_str.split("\n"), key=len)

            if len(val_str) > max_len:

                max_len = len(val_str)

        col_letter = openpyxl.utils.get_column_letter(col[0].column)

        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 40)

    file_stream = BytesIO()

    wb.save(file_stream)

    file_stream.seek(0)

    return StreamingResponse(

        file_stream,

        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",

        headers={"Content-Disposition": "attachment; filename=Production_Compound_Requirement.xlsx"}

    )

FRONTEND_DIST_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "frontend", "dist")

os.makedirs(FRONTEND_DIST_DIR, exist_ok=True)

app.mount("/", StaticFiles(directory=FRONTEND_DIST_DIR, html=True), name="static")
