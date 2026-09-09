import openpyxl

import os

import shutil

import re

import sqlite3

from .database import get_db_connection

def set_cell_value(sheet, row, col, value):

    # Find top-left cell if coordinate is inside a merged range

    for r in sheet.merged_cells.ranges:

        min_col, min_row, max_col, max_row = r.bounds

        if min_row <= row <= max_row and min_col <= col <= max_col:

            sheet.cell(row=min_row, column=min_col).value = value

            return

    # Direct write for non-merged cells

    sheet.cell(row=row, column=col).value = value

def clean_sheet_name(name):

    clean = re.sub(r'[:\\/\?\*\[\]]', '', name)

    return clean[:31]

def get_product_groups(sheet):

    groups = {}

    current_item_name = None

    current_sr_no = None

    for r in range(3, sheet.max_row + 1):

        sr_val = sheet.cell(row=r, column=1).value

        item_val = sheet.cell(row=r, column=2).value

        compound_val = sheet.cell(row=r, column=3).value

        side_val = sheet.cell(row=r, column=4).value

        if item_val is not None:

            current_item_name = str(item_val).strip()

            current_sr_no = sr_val

        if not current_item_name:

            continue

        if compound_val is None and side_val is None:

            row_vals = [sheet.cell(row=r, column=c).value for c in range(1, 16)]

            if not any(row_vals):

                continue

        thou_val = sheet.cell(row=r, column=5).value

        die_type = sheet.cell(row=r, column=6).value

        die_name = sheet.cell(row=r, column=7).value

        die_size = sheet.cell(row=r, column=8).value

        sheet_size = sheet.cell(row=r, column=9).value

        per_sheet_item = sheet.cell(row=r, column=10).value

        per_sheet_gm = sheet.cell(row=r, column=13).value

        row_data = {

            "row_index": r,

            "sr_no": current_sr_no,

            "item_name": current_item_name,

            "compound": str(compound_val).strip() if compound_val else None,

            "side": str(side_val).strip() if side_val else None,

            "thou": thou_val,

            "die_type": str(die_type).strip() if die_type else None,

            "die_name": str(die_name).strip() if die_name else None,

            "die_size": str(die_size).strip() if die_size else None,

            "comp_sheet_size": str(sheet_size).strip() if sheet_size else None,

            "per_sheet_item": per_sheet_item,

            "per_sheet_gm": per_sheet_gm

        }

        if current_item_name not in groups:

            groups[current_item_name] = []

        groups[current_item_name].append(row_data)

    return groups

def fuzzy_match_product(pi_prod_desc, master_keys):

    desc_clean = str(pi_prod_desc).strip().upper()

    for key in master_keys:

        if key.upper() == desc_clean:

            return key

    desc_no_space = desc_clean.replace(" ", "")

    for key in master_keys:

        if key.upper().replace(" ", "") == desc_no_space:

            return key

    match = re.search(r'([A-Z]{2,}\s*\d+)', desc_clean)

    if match:

        term = match.group(1).replace(" ", "")

        term_norm = re.sub(r'0+(\d+)', r'\1', term)

        for key in master_keys:

            key_clean = key.upper().replace(" ", "")

            key_norm = re.sub(r'0+(\d+)', r'\1', key_clean)

            if term_norm in key_norm or key_norm in term_norm:

                return key

    for key in master_keys:

        if key.upper() in desc_clean or desc_clean in key.upper():

            return key

    desc_words = set(re.findall(r'[A-Z0-9]+', desc_clean))

    for key in master_keys:

        key_words = set(re.findall(r'[A-Z0-9]+', key.upper()))

        if len(desc_words.intersection(key_words)) >= 2:

            return key

    return None

def match_pi_items_to_master(customer_id, pi_items, conn, plan_filepath):

    wb = openpyxl.load_workbook(plan_filepath, data_only=True)

    cursor = conn.cursor()

    cursor.execute("SELECT name FROM customers WHERE id = ?", (customer_id,))

    cust_row = cursor.fetchone()

    master_sheet = wb["LVP + SVG"]

    if cust_row:

        cust_name = cust_row["name"]

        clean_tab_name = clean_sheet_name(cust_name)

        if clean_tab_name in wb.sheetnames:

            master_sheet = wb[clean_tab_name]

    master_groups = get_product_groups(master_sheet)

    master_keys = list(master_groups.keys())

    matched_results = []

    for pi_item in pi_items:

        desc = pi_item["product_desc"]

        code = pi_item["product_code"]

        qty = pi_item["quantity"]

        cursor.execute(

            "SELECT planning_item_name FROM product_mappings WHERE customer_id = ? AND pi_product_name = ?",

            (customer_id, desc)

        )

        row = cursor.fetchone()

        matched_item_name = None

        match_confidence = "low"

        if row:

            matched_item_name = row["planning_item_name"]

            match_confidence = "remembered"

        else:

            matched_item_name = fuzzy_match_product(desc, master_keys)

            if not matched_item_name and code:

                matched_item_name = fuzzy_match_product(code, master_keys)

            if matched_item_name:

                match_confidence = "high"

        group_details = []

        if matched_item_name and matched_item_name in master_groups:

            group_details = master_groups[matched_item_name]

        matched_results.append({

            "product_code": code,

            "product_desc": desc,

            "quantity": qty,

            "matched_planning_item": matched_item_name,

            "confidence": match_confidence,

            "components": group_details

        })

    return matched_results, master_keys

def calculate_sheet_values(comp_size_str, per_sheet_gm, per_sheet_item, order_qty):

    try:

        per_sheet_item = int(per_sheet_item) if per_sheet_item is not None else 0

    except:

        per_sheet_item = 0

    try:

        order_qty = int(order_qty) if order_qty is not None else 0

    except:

        order_qty = 0

    required_sheet = 0.0

    if per_sheet_item > 0:

        required_sheet = order_qty / per_sheet_item

    width = 0.0

    height = 0.0

    if comp_size_str:

        match = re.findall(r'(\d+\.?\d*)', str(comp_size_str).upper())

        if len(match) >= 2:

            try:

                width = float(match[0])

                height = float(match[1])

            except:

                pass

        elif len(match) == 1:

            try:

                width = float(match[0])

                height = float(match[0])

            except:

                pass

    try:

        per_sheet_gm = float(per_sheet_gm) if per_sheet_gm is not None else 0.0

    except:

        per_sheet_gm = 0.0

    per_sheet_weight_gm = width * height * per_sheet_gm

    total_kg = (per_sheet_weight_gm * required_sheet) / 1000.0

    return required_sheet, per_sheet_weight_gm, total_kg

def process_customer_planning_sheet(customer_name, confirmed_mappings, plan_filepath, conn, upload_id):

    backup_path = plan_filepath.replace(".xlsx", "_backup.xlsx")

    if not os.path.exists(backup_path):

        shutil.copyfile(plan_filepath, backup_path)

    cursor = conn.cursor()

    cursor.execute("SELECT id FROM customers WHERE name = ?", (customer_name,))

    cust_row = cursor.fetchone()

    if cust_row:

        customer_id = cust_row["id"]

    else:

        cursor.execute("INSERT INTO customers (name) VALUES (?)", (customer_name,))

        conn.commit()

        customer_id = cursor.lastrowid

    for mapping in confirmed_mappings:

        pi_desc = mapping["product_desc"]

        plan_item = mapping["matched_planning_item"]

        cursor.execute(

            "INSERT INTO product_mappings (customer_id, pi_product_name, planning_item_name) "

            "VALUES (?, ?, ?) "

            "ON CONFLICT(customer_id, pi_product_name) DO UPDATE SET planning_item_name = ?",

            (customer_id, pi_desc, plan_item, plan_item)

        )

    conn.commit()

    wb = openpyxl.load_workbook(plan_filepath, data_only=False)

    clean_tab_name = clean_sheet_name(customer_name)

    if clean_tab_name in wb.sheetnames:

        sheet = wb[clean_tab_name]

    else:
        sheet = safe_copy_sheet(wb, "LVP + SVG", clean_tab_name)

    for r in range(3, sheet.max_row + 1):

        set_cell_value(sheet, r, 11, None)

        set_cell_value(sheet, r, 12, None)

        set_cell_value(sheet, r, 14, None)

        set_cell_value(sheet, r, 15, None)

    wb_read = openpyxl.load_workbook(plan_filepath, data_only=True)

    read_sheet = wb_read[clean_tab_name] if clean_tab_name in wb_read.sheetnames else wb_read["LVP + SVG"]

    sheet_groups = get_product_groups(read_sheet)

    # Build normalized map of sheet groups to support space-insensitive matching

    normalized_sheet_groups = {}

    for key, val in sheet_groups.items():

        norm_key = str(key).upper().replace(" ", "").strip()

        normalized_sheet_groups[norm_key] = val

    records_to_save = []

    for mapping in confirmed_mappings:

        plan_item = mapping["matched_planning_item"]

        qty = mapping["quantity"]

        if not plan_item:

            continue

        norm_plan_item = str(plan_item).upper().replace(" ", "").strip()

        # Check direct normalized match or fallback to fuzzy match

        group_rows = None

        if norm_plan_item in normalized_sheet_groups:

            group_rows = normalized_sheet_groups[norm_plan_item]

        else:

            matched_key = fuzzy_match_product(plan_item, sheet_groups.keys())

            if matched_key:

                norm_matched = str(matched_key).upper().replace(" ", "").strip()

                if norm_matched in normalized_sheet_groups:

                    group_rows = normalized_sheet_groups[norm_matched]

        if not group_rows:

            continue

        for row_info in group_rows:

            r = row_info["row_index"]

            gm = row_info["per_sheet_gm"]

            sz = row_info["comp_sheet_size"]

            is_required = True

            if gm is None or sz is None:

                is_required = False

            else:

                gm_str = str(gm).upper().strip()

                sz_str = str(sz).upper().strip()

                if "N/A" in gm_str or "VALUE" in gm_str or "DIV" in gm_str or not sz_str or sz_str == "NONE":

                    is_required = False

            if is_required:

                set_cell_value(sheet, r, 11, qty)

                required_sheet, per_sheet_weight, total_kg = calculate_sheet_values(

                    row_info["comp_sheet_size"],

                    row_info["per_sheet_gm"],

                    row_info["per_sheet_item"],

                    qty

                )

                set_cell_value(sheet, r, 12, f"=K{r}/J{r}")

                set_cell_value(sheet, r, 14, f'=LEFT(I{r},FIND("X",I{r})-1)*RIGHT(I{r},LEN(I{r})-FIND("X",I{r})-1)*M{r}')

                set_cell_value(sheet, r, 15, f"=(N{r}*L{r})/1000")

                records_to_save.append((

                    customer_id, upload_id, row_info["item_name"], row_info["compound"],

                    row_info["side"], row_info["thou"], row_info["die_type"], row_info["die_name"],

                    row_info["die_size"], row_info["comp_sheet_size"], row_info["per_sheet_item"],

                    qty, required_sheet, row_info["per_sheet_gm"], per_sheet_weight, total_kg

                ))

            else:

                set_cell_value(sheet, r, 11, None)

                set_cell_value(sheet, r, 12, None)

                set_cell_value(sheet, r, 14, None)

                set_cell_value(sheet, r, 15, None)

    wb.save(plan_filepath)

    cursor.execute("DELETE FROM planning_records WHERE upload_id = ?", (upload_id,))

    if records_to_save:

        cursor.executemany(

            "INSERT INTO planning_records (customer_id, upload_id, item_name, compound, side, thou, die_type, "

            "die_name, die_size, comp_sheet_size, per_sheet_item, order_qty, required_sheet, per_sheet_gm, "

            "per_sheet_weight_gm, total_kg) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",

            records_to_save

        )

    conn.commit()

    return len(records_to_save)

def safe_copy_sheet(wb, source_name, target_name):

    from copy import copy

    if target_name in wb.sheetnames:

        return wb[target_name]

    source = wb[source_name]

    target = wb.create_sheet(title=target_name)

    # Copy merged cells

    for merged_range in source.merged_cells.ranges:

        target.merge_cells(str(merged_range))

    # Copy cell values and styles

    for row in source.iter_rows():

        for cell in row:

            target_cell = target.cell(row=cell.row, column=cell.column, value=cell.value)

            if cell.has_style:

                target_cell.font = copy(cell.font)

                target_cell.fill = copy(cell.fill)

                target_cell.alignment = copy(cell.alignment)

                target_cell.border = copy(cell.border)

                target_cell.number_format = cell.number_format

    # Copy column widths

    for col_letter, col_dim in source.column_dimensions.items():

        if col_dim.width is not None:

            target.column_dimensions[col_letter].width = col_dim.width

    # Copy row heights

    for row_idx, row_dim in source.row_dimensions.items():

        if row_dim.height is not None:

            target.row_dimensions[row_idx].height = row_dim.height

    return target
